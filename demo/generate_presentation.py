import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: Modern DevSecOps Dark Theme
    COLOR_BG = RGBColor(15, 23, 42)       # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)     # Slate 800
    COLOR_PRIMARY = RGBColor(99, 102, 241) # Indigo 500
    COLOR_CYAN = RGBColor(6, 182, 212)    # Cyan 500
    COLOR_TEXT = RGBColor(248, 250, 252)  # Slate 50
    COLOR_MUTED = RGBColor(148, 163, 184) # Slate 400
    COLOR_ACCENT = RGBColor(16, 185, 129) # Emerald 500

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="DEVSECOPS CAPSTONE PROJECT"):
        # Category / eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(12)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ─────────────────────────────────────────────────────────────────────────
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide1)

    t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "CAMPUS RESOURCE HUB"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN

    p1 = tf1.add_paragraph()
    p1.text = "Enterprise-Grade DevSecOps Pipeline & Cloud-Native Deployment"
    p1.font.size = Pt(24)
    p1.font.color.rgb = COLOR_PRIMARY
    p1.space_before = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Automated CI/CD • Trivy Scanning • Terraform AWS EKS • Helm • ArgoCD GitOps • Prometheus & Grafana"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_before = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Author: Rohan Singh | GitHub: https://github.com/RohanSingh0208/DevOps-Project"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_ACCENT
    p3.space_before = Pt(30)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement & Application Domain
    # ─────────────────────────────────────────────────────────────────────────
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide2)
    add_header(slide2, "1. Project Overview & Application Architecture")

    # Left Column: Problem & Solution
    left_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf2_l = left_box.text_frame
    tf2_l.word_wrap = True

    p = tf2_l.paragraphs[0]
    p.text = "Campus Resource Hub Problem Context"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets_l = [
        ("Disorganized Campus Facilities", "Colleges struggle with conflicting paper/spreadsheet reservations for labs, lecture halls, and specialized equipment."),
        ("Unresolved Maintenance Deficits", "Broken projectors and faulty equipment lack structured reporting, tracking, and priority triage."),
        ("Full-Stack Solution", "Built an asynchronous FastAPI application paired with a modern React SPA interface to manage facility booking and maintenance lifecycles."),
    ]
    for title, desc in bullets_l:
        p_b = tf2_l.add_paragraph()
        p_b.text = f"• {title}: {desc}"
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT
        p_b.space_before = Pt(12)

    # Right Column: Technology Stack Matrix
    right_box = slide2.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf2_r = right_box.text_frame
    tf2_r.word_wrap = True

    p = tf2_r.paragraphs[0]
    p.text = "Technology Stack Breakdown"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    stack = [
        ("Backend", "Python 3.12, FastAPI, SQLAlchemy (Async/Sync), Alembic Migrations, Pydantic v2"),
        ("Frontend", "React 18, Vite, TypeScript, Lucide Icons, Recharts Analytics, Nginx Reverse Proxy"),
        ("Database", "PostgreSQL 16 with persistent volume and schema version control"),
        ("Cloud & Infra", "Terraform 1.5+, AWS VPC (ap-south-1), Amazon EKS 1.30 Managed Node Groups"),
        ("Orchestration", "Kubernetes, Helm 3, Ingress-Nginx, Horizontal Pod Autoscaler (HPA)"),
        ("GitOps & SRE", "ArgoCD, Prometheus Operator, ServiceMonitor, Grafana Observability"),
    ]
    for comp, tech in stack:
        p_s = tf2_r.add_paragraph()
        p_s.text = f"• {comp}: {tech}"
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = COLOR_TEXT
        p_s.space_before = Pt(8)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 3: Quality Engineering & Automated Testing
    # ─────────────────────────────────────────────────────────────────────────
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide3)
    add_header(slide3, "2. Testing, Code Quality & Database Migrations")

    box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf3 = box3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "Deterministic, Isolated Testing Strategy (29 Test Cases - 100% Passing)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    test_items = [
        ("Comprehensive Pytest Suite", "29 automated test cases covering /health, /ready, /metrics, CRUD operations for Resources, Reservations, Maintenance, and Dashboard stats."),
        ("Strict Overlap Validation", "Rigorous validation prevents double-booking of any facility. End times must strictly follow start times, returning 409 Conflict on collision."),
        ("Isolated Test Harness", "Utilizes SQLite in-memory fixtures (aiosqlite) with session overrides. Zero dependency on AWS or active production databases during test runs."),
        ("Database Migrations via Alembic", "Declarative schema revisions (0001_initial_schema.py) ensure reproducible table generation across dev, test, and production environments."),
        ("Code Quality & Linting", "Ruff linter integration enforces PEP8 compliance, import ordering, and type integrity across all service modules."),
    ]
    for title, desc in test_items:
        p_t = tf3.add_paragraph()
        p_t.text = f"✔ {title}: {desc}"
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_TEXT
        p_t.space_before = Pt(14)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 4: Containerization & DevSecOps Security
    # ─────────────────────────────────────────────────────────────────────────
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide4)
    add_header(slide4, "3. Container Architecture & Trivy Vulnerability Gates")

    box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf4 = box4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "Multi-Stage Slim Builds & Automated Security Scanning"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    sec_items = [
        ("Backend Containerization", "Multi-stage Python 3.12-slim build. C-extensions compiled in builder stage and stripped in final image. Runs as non-root user 'appuser' (UID 1001)."),
        ("Frontend Containerization", "Two-stage Node 20 alpine compilation to production static bundle, served by 'nginxinc/nginx-unprivileged:1.27-alpine' (UID 101) on port 8080."),
        ("Trivy Vulnerability Scanning", "Automated Aqua Security Trivy scan embedded directly in CI/CD pipeline with strict security gate: --severity HIGH,CRITICAL --exit-code 1."),
        ("Safe Migration Lifecycle", "Alembic migrations executed via dedicated one-shot service in Docker Compose / Kubernetes job to prevent multi-replica startup race conditions."),
        ("Immutable SHA Tagging", "Images published to GitHub Container Registry (GHCR) with immutable Git commit SHAs, guaranteeing traceability from source commit to production container."),
    ]
    for title, desc in sec_items:
        p_s = tf4.add_paragraph()
        p_s.text = f"🛡 {title}: {desc}"
        p_s.font.size = Pt(14)
        p_s.font.color.rgb = COLOR_TEXT
        p_s.space_before = Pt(14)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 5: Infrastructure as Code: Terraform & AWS
    # ─────────────────────────────────────────────────────────────────────────
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide5)
    add_header(slide5, "4. Infrastructure as Code: AWS VPC & Amazon EKS")

    box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf5 = box5.text_frame
    tf5.word_wrap = True

    p = tf5.paragraphs[0]
    p.text = "Modular, Cost-Conscious Cloud Infrastructure (ap-south-1 Mumbai)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    tf_items = [
        ("Dedicated VPC Architecture", "VPC CIDR 10.0.0.0/16 with DNS hostnames and resolution enabled. 2 public subnets (10.0.1.0/24, 10.0.2.0/24) and 2 private subnets (10.0.11.0/24, 10.0.12.0/24)."),
        ("Gateway & Routing Configuration", "Internet Gateway attached for public ingress; Elastic IP with NAT Gateway handles secure private worker node internet egress."),
        ("Amazon EKS 1.30 Control Plane", "Managed Kubernetes control plane spanning public and private endpoints with AmazonEKSClusterPolicy IAM roles."),
        ("Managed Worker Node Group", "Cost-conscious 2x t3.medium EC2 instances with autoscaling bounds (1 min, 4 max), attached to private subnets with least-privilege IAM policies."),
        ("OIDC Identity Provider (IRSA)", "IAM Roles for Service Accounts configured via TLS thumbprint, enabling pods to access AWS services securely without long-lived credentials."),
    ]
    for title, desc in tf_items:
        p_tf = tf5.add_paragraph()
        p_tf.text = f"☁ {title}: {desc}"
        p_tf.font.size = Pt(14)
        p_tf.font.color.rgb = COLOR_TEXT
        p_tf.space_before = Pt(14)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 6: Kubernetes Orchestration & Helm Packaging
    # ─────────────────────────────────────────────────────────────────────────
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide6)
    add_header(slide6, "5. Kubernetes Workloads, Helm Packaging & HPA")

    box6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf6 = box6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "Production-Ready Helm Chart Architecture (campus-resource-hub)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    k8s_items = [
        ("High Availability Deployments", "2 replicas each for Frontend and Backend in production (values.yaml), with a lightweight 1-replica profile for local dev (values-dev.yaml)."),
        ("Ingress-Nginx Routing", "Single unified ingress resource routing '/' traffic to Frontend SPA and '/api' + '/metrics' directly to Backend service on port 8000."),
        ("Horizontal Pod Autoscaler (HPA)", "Autoscales backend pods between 2 and 6 replicas when CPU utilization exceeds 70%, with stabilization windows preventing thrashing."),
        ("Liveness & Readiness Probes", "Strict HTTP probes (/healthz on frontend, /health and /ready DB probe on backend) ensure zero-downtime rolling updates."),
        ("Stateful Persistence", "PostgreSQL deployed with dedicated PersistentVolumeClaim (5Gi) ensuring transactional data survives pod restarts."),
    ]
    for title, desc in k8s_items:
        p_k = tf6.add_paragraph()
        p_k.text = f"☸ {title}: {desc}"
        p_k.font.size = Pt(14)
        p_k.font.color.rgb = COLOR_TEXT
        p_k.space_before = Pt(14)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 7: GitOps Continuous Delivery with ArgoCD
    # ─────────────────────────────────────────────────────────────────────────
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide7)
    add_header(slide7, "6. Continuous Delivery: The ArgoCD GitOps Loop")

    box7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf7 = box7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "Declarative, Automated GitOps Workflow"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    gitops_items = [
        ("Git as Single Source of Truth", "Repository at https://github.com/RohanSingh0208/DevOps-Project holds all manifests, Helm templates, and environment values."),
        ("Automated CI Pipeline Trigger", "GitHub Actions tests code, builds images, scans with Trivy, publishes to GHCR, and commits updated image SHA back to values.yaml."),
        ("ArgoCD Reconciliation Loop", "In-cluster ArgoCD controller continuously polls the repository, detecting differences between desired Git state and live cluster state."),
        ("Self-Healing & Auto-Sync", "ArgoCD automatically remediates configuration drift: manual changes made to cluster resources are immediately overwritten by Git truth."),
        ("Auditability & Instant Rollback", "Every deployment is bound to an exact Git commit. Rolling back production requires nothing more than 'git revert'."),
    ]
    for title, desc in gitops_items:
        p_g = tf7.add_paragraph()
        p_g.text = f"🔄 {title}: {desc}"
        p_g.font.size = Pt(14)
        p_g.font.color.rgb = COLOR_TEXT
        p_g.space_before = Pt(14)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 8: Monitoring & Troubleshooting Lab
    # ─────────────────────────────────────────────────────────────────────────
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide8)
    add_header(slide8, "7. Observability (Prometheus/Grafana) & Troubleshooting")

    box8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.2))
    tf8 = box8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "Full Stack Observability & Real-World Diagnostic Defense"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    obs_items = [
        ("Prometheus Operator & ServiceMonitor", "FastAPI instrumentator automatically exposes RED metrics at /metrics. ServiceMonitor scrapes backend pods every 15 seconds."),
        ("Production Grafana Dashboard", "Custom dashboard monitoring Request Rates, p95 Latency, 4xx/5xx Error Rates, Pod CPU/Memory, and Ready Replica counts."),
        ("Lab Scenario 1: ImagePullBackOff", "Deliberately broken image tag manifest diagnosed via 'kubectl describe pod' and events, demonstrating image registry troubleshooting."),
        ("Lab Scenario 2: Service Selector Mismatch", "Deliberately mismatched selector resulting in empty endpoints (<none>), diagnosed via 'kubectl get endpoints' and repaired with kubectl patch."),
        ("Key Engineering Takeaways", "DevSecOps is not just tools; it is the discipline of automated security gates, reproducible infrastructure, and observability-driven verification."),
    ]
    for title, desc in obs_items:
        p_o = tf8.add_paragraph()
        p_o.text = f"📊 {title}: {desc}"
        p_o.font.size = Pt(14)
        p_o.font.color.rgb = COLOR_TEXT
        p_o.space_before = Pt(14)

    os.makedirs("demo", exist_ok=True)
    output_path = os.path.join("demo", "presentation.pptx")
    prs.save(output_path)
    print(f"[OK] PowerPoint presentation successfully generated at: {output_path}")

if __name__ == "__main__":
    create_presentation()
